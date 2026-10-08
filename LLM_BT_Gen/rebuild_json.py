# -*- coding: utf-8 -*-
"""Rebuild missing summary_<...>_<N>.json files from their xlsx counterpart.

The stats scripts (stat_avg_all.py / stats_avg.py) read ONLY the json, but
some rounds lost their json while the xlsx survived.  This script reconstructs
a stats-compatible json from the xlsx rows:
  * one record per xlsx row (task, success, calls, times, tokens, cost)
  * success=0 rows with zero LLM calls get an 'LLM error' reason so that
    repair_task.py treats them as network failures to re-run
  * success=0 rows WITH calls get 'not verified' (a genuine model failure)
  * benchmark tasks with NO row at all (row deleted from the xlsx) become
    placeholder 'LLM error: no record' records -- repair_task.py re-runs
    them and appends the row
Existing json files are never overwritten (--force to rebuild anyway).

Usage:
  python rebuild_json.py                    # results_complex defaults
  python rebuild_json.py --results_dir results_hard \
      --benchmark ../behaverify/Benchmark/Behaverify_benchmark_hard.xlsx
"""
import argparse
import glob
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))


def benchmark_tasks(path):
    """Task names from the benchmark xlsx ('benchmark' sheet, column A)."""
    import openpyxl
    ws = openpyxl.load_workbook(path)['benchmark']
    return [str(r[0]).strip() for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]


def placeholder(task):
    return {'task': task, 'success': False,
            'reason': 'LLM error: no record (row deleted from xlsx)',
            'final_tree': None, 'rounds': 0, 'n_calls': 0,
            'time_s': 0, 'llm_time_s': 0, 'verify_time_s': 0,
            'prompt_tokens': 0, 'completion_tokens': 0, 'total_tokens': 0,
            'cost_usd': 0}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--results_dir', default=os.path.join(HERE, 'results_complex'))
    ap.add_argument('--benchmark', default=os.path.join(HERE, 'complex_tasks.xlsx'))
    ap.add_argument('--force', action='store_true',
                    help='rebuild even when the json already exists')
    args = ap.parse_args()

    bench = benchmark_tasks(args.benchmark)
    for xf in sorted(glob.glob(os.path.join(args.results_dir, 'summary_*.xlsx'))):
        stem = os.path.basename(xf)[:-5]
        m = re.match(r'summary_(.*)_(\d+)$', stem)
        if not m:
            print('skip (unparsed name):', stem)
            continue
        model, baseline = m.group(1), 'full'
        for suf in ('_noce', '_withbt'):
            if model.endswith(suf):
                model, baseline = model[:-len(suf)], suf[1:]
        jpath = xf[:-5] + '.json'
        if os.path.exists(jpath) and not args.force:
            continue

        import openpyxl
        rows = [r for r in openpyxl.load_workbook(xf).active.iter_rows(
            min_row=2, values_only=True) if r[0]]
        recs = []
        for r in rows:
            rec = {'task': str(r[0]), 'success': bool(r[1]), 'final_tree': None,
                   'n_calls': r[3] or 0, 'rounds': r[3] or 0,
                   'time_s': r[4], 'llm_time_s': r[5], 'verify_time_s': r[6],
                   'prompt_tokens': r[7] or 0, 'completion_tokens': r[8] or 0,
                   'total_tokens': r[9] or 0, 'cost_usd': r[10]}
            if not rec['success']:
                rec['reason'] = ('LLM error: rebuilt from xlsx, zero LLM calls'
                                 if not rec['n_calls'] else
                                 'not verified within 5 rounds (rebuilt from xlsx)')
            recs.append(rec)
        have = {r['task'] for r in recs}
        recs.extend(placeholder(t) for t in bench if t not in have)
        with open(jpath, 'w', encoding='utf-8') as f:
            json.dump({'model': model, 'baseline': baseline, 'max_rounds': 5,
                       'results': recs}, f, ensure_ascii=False, indent=1)
        print('rebuilt {}  ({} rows, {} placeholders)'.format(
            os.path.basename(jpath), len(rows), len(bench) - len(rows)))


if __name__ == '__main__':
    main()
