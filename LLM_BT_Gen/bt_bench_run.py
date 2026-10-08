# -*- coding: utf-8 -*-
"""BTRun -- portable benchmark runner for LLM-driven BehaVerify BT synthesis.

Packs the whole experiment (hard 70-task benchmark + complex 10-task
benchmark, 3 models x 3 baselines, CEGIS loop with behaverify/NuSMV
verification) into ONE portable executable (Windows: BTRun.exe,
Linux/Ubuntu 22.04+: BTRun).  Copy the distribution folder to any machine
and run it -- no Python, no venv, no behaverify checkout and no NuSMV
installation is required on the target machine.

Modes (selected by first flags):
  (no flag)      orchestration mode: fans out worker processes, one per
                 model (--parallel models, default) or per baseline
                 (--parallel baselines), or runs everything sequentially
                 (--parallel none).  Worker stdout is tagged and echoed.
  --worker       internal: run one list of jobs sequentially (never call
                 this by hand; used by the orchestrator).
  --verify-child internal: run behaverify nuxmv + NuSMV on one .tree file
                 inside this process (replaces `python -m behaverify`).
  --selftest     check the bundled toolchain (NuSMV + behaverify verify on a
                 reference tree); add --api to also test the LLM connection.

Network policy: the LLM API is always contacted DIRECTLY (no proxy).  A
"proxy" entry in an external llm_config.json is deliberately ignored; the
only way to enable a proxy is the LLM_PROXY environment variable.

All output and all generated files are plain ASCII/English.

Usage examples:
  BTRun.exe                                   # hard + complex, 3 models x
                                              # 3 baselines x 5 repeats,
                                              # parallel across models
  BTRun.exe --suite hard --repeat 1           # quick hard-only pilot
  BTRun.exe --parallel baselines              # parallel across baselines
  BTRun.exe --suite complex --models gpt-4o --baselines full --repeat 5
  BTRun.exe --limit 2 --repeat 1              # smoke test, 2 tasks/suite
  BTRun.exe --selftest                        # toolchain self check
  BTRun.exe --selftest --api                  # + live LLM connectivity check
"""
import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time

# ---------------------------------------------------------------- app layout
FROZEN = getattr(sys, 'frozen', False)
if os.environ.get('BT_APP_DIR'):
    APP_DIR = os.path.abspath(os.environ['BT_APP_DIR'])   # explicit override
elif FROZEN:
    APP_DIR = os.path.dirname(os.path.abspath(sys.executable))
else:
    APP_DIR = os.path.dirname(os.path.abspath(__file__))

# NuSMV binary shipped in the distribution folder: NuSMV.exe on Windows,
# NuSMV on Linux -- both under NuSMV/bin/.
NUSMV_BIN = 'NuSMV.exe' if os.name == 'nt' else 'NuSMV'
NUXMV = os.path.join(APP_DIR, 'NuSMV', 'bin', NUSMV_BIN)
BENCH_DIR = os.path.join(APP_DIR, 'benchmarks')
BENCHMARKS = {
    'hard': os.path.join(BENCH_DIR, 'Behaverify_benchmark_hard.xlsx'),
    'complex': os.path.join(BENCH_DIR, 'complex_tasks.xlsx'),
}
PROMPT_FILE = os.path.join(APP_DIR, 'prompt_behaverify.txt')
SELFTEST_TREE = os.path.join(APP_DIR, 'selftest', 'selftest.tree')

# ---- verification toolchain -------------------------------------------------
VERIFY_TIMEOUT = 180          # seconds per behaverify run

# ---- experiment grid ----------------------------------------------------------
MODELS = ('gpt-4o', 'gpt-5.5', 'claude-opus-4-6')
BASELINES = ('full', 'noce', 'withbt')

# ---- LLM connection (embedded defaults, direct connection, NO proxy) ---------
DEFAULT_BASE_URL = 'https://4sapi.org'
DEFAULT_KEYS = {
    'gpt': 'sk-Z36I8HHBjjEASKByE6xbSEQq76CHmPD20wCVDvWPGTnpaMmD',
    'claude': 'sk-RoXVfA3Fno4VPwxU7h6MgpDjz1NAJ4m8hBbrfXuPkLjBWGMV',
}

# USD per 1M tokens, {model prefix: [input, output]}, longest prefix wins.
DEFAULT_PRICES = {
    'gpt-4o-mini': [0.15, 0.60],
    'gpt-4o': [2.50, 10.00],
    'gpt-5.5': [5.00, 30.00],
    'claude-opus': [5.00, 25.00],
    'claude-sonnet': [3.00, 15.00],
    'claude-haiku': [1.00, 5.00],
    'gemini-2.5-pro': [1.25, 10.00],
    'gemini-2.5-flash': [0.30, 2.50],
}

# ---- LLM call retry policy ---------------------------------------------------
LLM_MAX_ATTEMPTS = 30         # attempts for transport-level errors
LLM_RETRY_WAIT = 1.0          # seconds between attempts (fixed)

SPEC_LINE_RE = re.compile(r'-- specification\s+(.+?)\s+is\s+(true|false)', re.IGNORECASE)


def _install_tkinter_stub():
    """The verification path never touches the GUI, but behaverify imports
    tkinter at module level (behaverify_gui).  The frozen build deliberately
    does not ship the Tk runtime, so if importing the real tkinter fails,
    install lightweight stubs that satisfy the imports."""
    try:
        import tkinter  # noqa: F401
        return
    except Exception:
        pass
    import types
    stubs = {}
    for name in ('tkinter', 'tkinter.ttk', 'tkinter.filedialog',
                 'tkinter.messagebox', 'tkinter.scrolledtext',
                 'tkinter.font', 'tkinter.colorchooser', 'tkinter.simpledialog'):
        mod = types.ModuleType(name)
        mod.__path__ = []
        stubs[name] = mod
    stubs['tkinter'].ttk = stubs['tkinter.ttk']
    stubs['tkinter'].filedialog = stubs['tkinter.filedialog']
    stubs['tkinter'].messagebox = stubs['tkinter.messagebox']
    stubs['tkinter'].scrolledtext = stubs['tkinter.scrolledtext']
    stubs['tkinter'].font = stubs['tkinter.font']
    for name, mod in stubs.items():
        sys.modules[name] = mod


def self_exe_cmd():
    """Command prefix to relaunch this program (exe when frozen, python+file
    in development)."""
    if FROZEN:
        return [sys.executable]
    return [sys.executable, os.path.abspath(__file__)]


def child_env(extra=None):
    """Environment for child processes: UTF-8, unbuffered, NuSMV on PATH."""
    env = dict(os.environ)
    env['PYTHONIOENCODING'] = 'utf-8'
    env['PYTHONUNBUFFERED'] = '1'
    nuxmv_bin = os.path.dirname(NUXMV)
    if os.path.isdir(nuxmv_bin):
        env['PATH'] = nuxmv_bin + os.pathsep + env.get('PATH', '')
        # Linux: bundled NuSMV shared libraries (if any) next to the binary
        nuxmv_lib = os.path.join(os.path.dirname(nuxmv_bin), 'lib')
        if os.name != 'nt' and os.path.isdir(nuxmv_lib):
            env['LD_LIBRARY_PATH'] = (nuxmv_lib + os.pathsep
                                      + env.get('LD_LIBRARY_PATH', ''))
    if extra:
        env.update(extra)
    return env


# ---------------------------------------------------------------- benchmark IO
def load_benchmark(xlsx_path):
    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path)
    ws = wb[wb.sheetnames[0]]
    tasks = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        task, desc, ltl = row[0], row[1], row[2]
        if not task:
            continue
        specs = [ln.strip() for ln in str(ltl or '').splitlines() if ln.strip()]
        tasks.append({'task': str(task).strip(), 'description': str(desc or '').strip(),
                      'specs': specs})
    return tasks


# ---------------------------------------------------------------- results xlsx
XLSX_HEADERS = ['Task', 'Success', 'BT', 'LLM Calls', 'Time (s)',
                'LLM Time (s)', 'Verify Time (s)',
                'Prompt Tokens', 'Completion Tokens', 'Total Tokens',
                'Cost (USD)']


def next_summary_paths(results_dir, model, baseline='full'):
    """Per-run summary paths: summary_<model>_<N> for the full baseline,
    summary_<model>_<baseline>_<N> otherwise.  N is the first free number so
    repeated runs stack up as _1, _2, _3 ..."""
    stem = 'summary_' + re.sub(r'[^\w.-]', '_', model)
    if baseline != 'full':
        stem += '_' + baseline
    n = 1
    while (os.path.exists(os.path.join(results_dir, stem + '_{}.json'.format(n)))
           or os.path.exists(os.path.join(results_dir, stem + '_{}.xlsx'.format(n)))):
        n += 1
    base = os.path.join(results_dir, '{}_{}'.format(stem, n))
    return base + '.json', base + '.xlsx'


def write_results_xlsx(summary, xlsx_path):
    """Per-task results xlsx (same layout as llm_bt_gen.py, rows matched by
    task name so separate runs accumulate in the same file)."""
    import openpyxl
    if os.path.exists(xlsx_path):
        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb.active
        for col, h in enumerate(XLSX_HEADERS, start=1):
            if not ws.cell(row=1, column=col).value:
                ws.cell(row=1, column=col, value=h)
        row_of = {ws.cell(row=r, column=1).value: r
                  for r in range(2, ws.max_row + 1)}
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(XLSX_HEADERS)
        row_of = {}
    for r in summary:
        bt = ''
        if r['success'] and r.get('final_tree') and os.path.exists(r['final_tree']):
            with open(r['final_tree'], encoding='utf-8') as f:
                bt = f.read()
        vals = [r['task'], 1 if r['success'] else 0, bt,
                r.get('n_calls'), r.get('time_s'),
                r.get('llm_time_s'), r.get('verify_time_s'),
                r.get('prompt_tokens'), r.get('completion_tokens'),
                r.get('total_tokens'), r.get('cost_usd')]
        if r['task'] in row_of:
            for col, v in enumerate(vals, start=1):
                ws.cell(row=row_of[r['task']], column=col, value=v)
        else:
            ws.append(vals)
    wb.save(xlsx_path)


# ---------------------------------------------------------------- LLM calling
def load_config():
    """Embedded defaults; an optional llm_config.json next to the executable
    may override base_url/model/keys/prices.  Any 'proxy' entry there is
    IGNORED on purpose -- this build always connects directly unless the
    LLM_PROXY environment variable is set explicitly."""
    cfg = {'base_url': DEFAULT_BASE_URL, 'proxy': '', 'model': 'gpt-4o',
           'keys': dict(DEFAULT_KEYS), 'prices': dict(DEFAULT_PRICES)}
    ext = os.path.join(APP_DIR, 'llm_config.json')
    if os.path.exists(ext):
        try:
            with open(ext, encoding='utf-8') as f:
                data = json.load(f)
            for k in ('base_url', 'model'):
                if data.get(k):
                    cfg[k] = data[k]
            if data.get('keys'):
                cfg['keys'].update({k: v for k, v in data['keys'].items() if v})
            if data.get('prices'):
                cfg['prices'].update(data['prices'])
            print('config: using overrides from llm_config.json (proxy entries ignored)')
        except Exception as e:
            print('config: could not read llm_config.json ({}), using built-in defaults'.format(e))
    cfg['model'] = os.environ.get('LLM_MODEL', cfg['model'])
    cfg['base_url'] = os.environ.get('LLM_BASE_URL', cfg['base_url'])
    cfg['proxy'] = os.environ.get('LLM_PROXY', '')   # empty = DIRECT connection
    select_api_key(cfg)
    return cfg


def select_api_key(cfg):
    """Pick the API key for cfg['model']: explicit env var wins, else the
    first matching model-name prefix from the key table."""
    cfg['api_key'] = os.environ.get('LLM_API_KEY', '')
    if not cfg['api_key']:
        for prefix, k in (cfg.get('keys') or {}).items():
            if cfg['model'].startswith(prefix):
                cfg['api_key'] = k
                break


# ---------------------------------------------------------------- cost model
def price_for(cfg, model=None):
    model = model or cfg['model']
    prices = dict(DEFAULT_PRICES)
    prices.update(cfg.get('prices') or {})
    for prefix in sorted(prices, key=len, reverse=True):
        if model.startswith(prefix):
            return prices[prefix]
    return None


def cost_of(cfg, tokens):
    price = price_for(cfg)
    if not price:
        return None
    return round(tokens['prompt_tokens'] / 1e6 * price[0]
                 + tokens['completion_tokens'] / 1e6 * price[1], 6)


def call_llm(cfg, prompt):
    """OpenAI-compatible chat call via http.client with a RAW Authorization
    header.  Direct HTTPS connection by default; only uses a CONNECT tunnel
    when LLM_PROXY is explicitly set (never reads the system proxy)."""
    import http.client
    from urllib.parse import urlparse
    host = urlparse(cfg['base_url']).hostname or '4sapi.org'
    payload = json.dumps({'model': cfg['model'],
                          'messages': [{'role': 'user', 'content': prompt}]})
    headers = {'Accept': 'application/json',
               'Authorization': cfg['api_key'],      # raw key, no Bearer prefix
               'Content-Type': 'application/json'}
    proxy = cfg.get('proxy') or ''
    if proxy:
        p = urlparse(proxy if '//' in proxy else 'http://' + proxy)
        conn = http.client.HTTPSConnection(p.hostname, p.port or 80, timeout=300)
        conn.set_tunnel(host, 443)
    else:
        conn = http.client.HTTPSConnection(host, timeout=300)
    conn.request('POST', '/v1/chat/completions', payload, headers)
    res = conn.getresponse()
    raw = res.read().decode('utf-8', errors='replace')
    if res.status != 200:
        raise RuntimeError('LLM API error {}: {}'.format(res.status, raw[:300]))
    data = json.loads(raw)
    content = data['choices'][0]['message']['content'] or ''
    return content, (data.get('usage') or {})


def call_llm_retry(cfg, prompt, log=print):
    """call_llm with retries for transient errors (SSL, timeouts, 5xx, ...).
    Only successful responses consume one of max_rounds.  Returns
    (reply, usage, llm_time_s); llm_time_s includes failed attempts."""
    max_attempts = cfg.get('llm_attempts') or LLM_MAX_ATTEMPTS
    retry_wait = cfg.get('llm_retry_wait')
    if retry_wait is None:
        retry_wait = LLM_RETRY_WAIT
    last, llm_time = None, 0.0
    for attempt in range(1, max_attempts + 1):
        t0 = time.time()
        try:
            reply, usage = call_llm(cfg, prompt)
            return reply, usage, llm_time + (time.time() - t0)
        except Exception as e:
            llm_time += time.time() - t0
            last = e
            if attempt < max_attempts:
                log('  LLM call failed (attempt {}/{}): {} -- retrying in {:.0f}s '
                    '(failed attempts do not count as LLM calls)'.format(
                        attempt, max_attempts, e, retry_wait))
                time.sleep(retry_wait)
    raise last


# ------------------------------------------------------------- tree extraction
def extract_tree(response_text):
    """Pull the .tree content out of the LLM reply (prefer a fenced block)."""
    m = re.search(r'```(?:\w*)\s*\n(.*?)```', response_text, re.DOTALL)
    text = m.group(1) if m else response_text
    start = text.find('configuration')
    if start == -1:
        return None
    return text[start:].strip()


def inject_specs(tree_text, specs):
    """Replace the specifications block with the canonical one from the xlsx."""
    block = 'specifications {\n' + '\n'.join('    ' + s for s in specs) + '\n}'
    m = re.search(r'specifications\s*\{.*?\}\s*$', tree_text, re.DOTALL)
    if m:
        return tree_text[:m.start()] + block + '\n'
    return tree_text.rstrip() + '\n' + block + '\n'


# ---------------------------------------------------------------- verification
def verify_tree(tree_path, out_dir):
    """Run behaverify generation + NuSMV LTL checking by relaunching this
    executable in --verify-child mode (behaverify runs in that process).
    Returns a report dict."""
    os.makedirs(out_dir, exist_ok=True)
    cmd = self_exe_cmd() + ['--verify-child', tree_path, out_dir,
                            '--generate', '--do_not_trim', '--ltl', '--overwrite']
    try:
        p = subprocess.run(cmd, capture_output=True, text=True,
                           encoding='utf-8', errors='replace',
                           timeout=VERIFY_TIMEOUT, env=child_env())
        rc, so, se = p.returncode, p.stdout or '', p.stderr or ''
    except subprocess.TimeoutExpired:
        return {'status': 'error', 'message': 'verification timed out ({}s)'.format(VERIFY_TIMEOUT),
                'results': [], 'n_true': 0, 'n_false': 0}
    name = os.path.splitext(os.path.basename(tree_path))[0]
    outfile = os.path.join(out_dir, 'nuxmv', name + '_output.txt')
    results = []
    if os.path.exists(outfile):
        with open(outfile, encoding='utf-8', errors='replace') as f:
            content = f.read()
        for m in SPEC_LINE_RE.finditer(content):
            results.append({'spec': m.group(1).strip(), 'result': m.group(2).lower()})
        counterexample = ''
        i = content.find('as demonstrated by')
        if i != -1:
            counterexample = content[i:i + 3000]
    else:
        content = ''
        counterexample = ''
    n_true = sum(1 for r in results if r['result'] == 'true')
    n_false = sum(1 for r in results if r['result'] == 'false')
    if not results:
        # generation or model-checking failed: surface the actual error text
        tail = (se + '\n' + so).strip().splitlines()
        err = '\n'.join(tail[-25:]) or ('no verification output produced (rc={})'.format(rc))
        return {'status': 'error', 'message': err, 'results': [], 'n_true': 0, 'n_false': 0}
    status = 'pass' if n_false == 0 else 'fail'
    return {'status': status, 'message': '', 'results': results,
            'n_true': n_true, 'n_false': n_false,
            'counterexample': counterexample}


def verify_child_main(tree_path, out_dir, rest_argv):
    """In-process `python -m behaverify nuxmv <tree> <out> ...` replacement."""
    _install_tkinter_stub()          # behaverify imports tkinter at module level
    from behaverify.behaverify import main as behaverify_main
    argv = ['behaverify', 'nuxmv', tree_path, out_dir] + rest_argv
    if '--nuxmv_path' not in rest_argv:
        argv += ['--nuxmv_path', NUXMV]
    sys.argv = argv
    behaverify_main()


def build_feedback(report, expected_specs):
    """Human/LLM-readable feedback for one failed round."""
    if report['status'] == 'error':
        return ('Your model could not be verified. The tool reported:\n'
                + report['message']
                + '\nFix the model so it is a valid BehaVerify .tree file whose '
                  'tree contains every node referenced by the specifications.')
    lines = []
    for r in report['results']:
        if r['result'] == 'false':
            lines.append('  FALSE: ' + r['spec'])
    fb = ('The verifier found {} violating specification(s) ({} of {} passed):\n'.format(
        len(lines), report['n_true'], len(report['results'])) + '\n'.join(lines))
    if report.get('counterexample'):
        fb += ('\n\nCounterexample trace (excerpt):\n' + report['counterexample'][:2500]
               + '\nThink about WHY the tree lets this trace happen (remember: '
                 'in a sequence, a failing child stops everything; in a selector, '
                 'a child that does NOT fail satisfies it) and produce a '
                 'corrected tree.')
    return fb


# --------------------------------------------------------------- one CEGIS run
def synthesize(cfg, task, max_rounds, baseline='full', results_dir='results',
               log=print):
    """Generate -> verify -> retry loop for one task.

    baseline decides what the LLM sees after a failed attempt:
      full   - previous attempt + verifier feedback (main method)
      withbt - previous attempt only, no verifier information
      noce   - nothing: the plain per-task prompt again (re-generation)
    """
    prompt_base = open(PROMPT_FILE, encoding='utf-8').read()
    # model-scoped task directory: parallel workers for DIFFERENT models may
    # run the SAME task at the same time, so the directory must contain the
    # model name or the workers would overwrite each other's round files
    model_dir = re.sub(r'[^\w.-]', '_', cfg['model'])
    task_dir = (os.path.join(results_dir, model_dir, task['task']) if baseline == 'full'
                else os.path.join(results_dir, model_dir, baseline, task['task']))
    os.makedirs(task_dir, exist_ok=True)
    history = []          # (round, attempt_text, retry_note) of the last round
    n_calls = 0           # successful LLM calls only
    tokens = {'prompt_tokens': 0, 'completion_tokens': 0, 'total_tokens': 0}
    llm_time = 0.0        # seconds spent inside call_llm (incl. failed attempts)
    verify_time = 0.0     # seconds spent inside verify_tree
    t0 = time.time()

    def done(**res):
        """Stamp task stats (calls / times / tokens / cost) onto a result."""
        res['task'] = task['task']
        res['n_calls'] = n_calls
        res['time_s'] = round(time.time() - t0, 1)
        res['llm_time_s'] = round(llm_time, 1)
        res['verify_time_s'] = round(verify_time, 1)
        res.update(tokens)
        res['cost_usd'] = cost_of(cfg, tokens)
        return res

    for rnd in range(1, max_rounds + 1):
        prompt = prompt_base + '\n\n' + '=' * 70 + '\n'
        prompt += 'TASK: {}\n'.format(task['task'])
        prompt += 'Instruction: {}\n'.format(task['description'])
        prompt += ('The specifications block of your file MUST contain exactly '
                   'these formulas (they will be inserted for you, structure the '
                   'tree so that all of them hold):\n'
                   + '\n'.join(task['specs']) + '\n')
        if history and baseline != 'noce':
            _, prev_attempt, note = history[-1]
            prompt += ('\n\n' + '=' * 70 + '\nYOUR PREVIOUS ATTEMPT:\n' + prev_attempt
                       + '\n\n' + note)
        if not history or baseline == 'noce':
            prompt += ('\nNow output the complete .tree file for the TASK above. '
                       'Output only the file content.')
        log('  round {}: calling LLM...'.format(rnd))
        try:
            reply, usage, call_s = call_llm_retry(cfg, prompt, log)
            n_calls += 1
            llm_time += call_s
            tokens['prompt_tokens'] += usage.get('prompt_tokens') or 0
            tokens['completion_tokens'] += usage.get('completion_tokens') or 0
            tokens['total_tokens'] += usage.get('total_tokens') or (
                (usage.get('prompt_tokens') or 0) + (usage.get('completion_tokens') or 0))
        except Exception as e:
            log('  LLM call still failing after retries: {}'.format(e))
            return done(success=False, rounds=rnd - 1,
                        reason='LLM error: {}'.format(e), final_tree=None)
        candidate = extract_tree(reply)
        if not candidate:
            note = ('Your previous reply did not contain a .tree model (no '
                    '"configuration" section found). Output only the complete '
                    'file content.')
            history.append((rnd, reply[:2000], note))
            continue
        candidate = inject_specs(candidate, task['specs'])
        cand_path = os.path.join(task_dir, 'round_{}.tree'.format(rnd))
        with open(cand_path, 'w', encoding='utf-8') as f:
            f.write(candidate)
        tv = time.time()
        report = verify_tree(cand_path, os.path.join(task_dir, 'round_{}_out'.format(rnd)))
        verify_time += time.time() - tv
        log('    verify: {} ({} true / {} false)'.format(
            report['status'], report['n_true'], report['n_false']))
        if report['status'] == 'error':
            log('      verify error detail: {}'.format(
                ' | '.join(report['message'].splitlines())[:400]))
        if report['status'] == 'pass':
            final = os.path.join(task_dir, 'final.tree')
            shutil.copy(cand_path, final)
            return done(success=True, rounds=rnd,
                        n_specs=len(report['results']), final_tree=final)
        if baseline == 'full':
            note = ('VERIFIER FEEDBACK:\n' + build_feedback(report, task['specs'])
                    + '\n\nProduce a corrected, complete .tree file.')
        else:
            # withbt / noce: keep the verifier output to ourselves
            note = ('The previous attempt did not satisfy the specifications. '
                    'Produce a corrected, complete .tree file.')
        history.append((rnd, candidate, note))
    return done(success=False, rounds=max_rounds,
                reason='not verified within {} rounds'.format(max_rounds),
                final_tree=None)


# ------------------------------------------------------------- worker: one job
def run_job(job, log):
    """One (suite x model x baseline x repeat) job: run the task set, write
    summary json + xlsx into the suite results directory."""
    cfg = load_config()
    cfg['model'] = job['model']
    cfg['llm_attempts'] = job.get('llm_attempts')
    cfg['llm_retry_wait'] = job.get('llm_retry_wait')
    select_api_key(cfg)
    if not cfg.get('api_key'):
        log('no API key for model {} -- job skipped'.format(job['model']))
        return 1
    tasks = load_benchmark(job['benchmark'])
    if job.get('tasks'):
        want = set(job['tasks'])
        tasks = [t for t in tasks if t['task'] in want]
    if job.get('limit'):
        tasks = tasks[:job['limit']]
    if not tasks:
        log('no matching tasks -- job skipped')
        return 1
    results_dir = job['results_dir']
    os.makedirs(results_dir, exist_ok=True)
    baseline = job['baseline']
    max_rounds = job.get('max_rounds') or 10
    repeat = job.get('repeat') or 1

    log('suite={} model={} baseline={} tasks={} repeat={} max_rounds={}'.format(
        job['suite'], job['model'], baseline, len(tasks), repeat, max_rounds))
    if price_for(cfg) is None:
        log('warning: no price configured for {} -- cost_usd will be null'.format(
            job['model']))

    for rep in range(1, repeat + 1):
        if repeat > 1:
            log('repeat {}/{}'.format(rep, repeat))
        summary = []
        t0 = time.time()
        for t in tasks:
            log('== {} =='.format(t['task']))
            r = synthesize(cfg, t, max_rounds, baseline=baseline,
                           results_dir=results_dir, log=log)
            cost = r.get('cost_usd')
            log('   -> {} (rounds={}, calls={}, {:.1f}s | llm {:.1f}s, verify '
                '{:.1f}s | tokens={} | cost={})'.format(
                    'SUCCESS' if r['success'] else 'FAILED', r['rounds'],
                    r['n_calls'], r['time_s'], r['llm_time_s'],
                    r['verify_time_s'], r['total_tokens'],
                    '${:.4f}'.format(cost) if cost is not None else 'n/a'))
            summary.append(r)
        ok = sum(1 for r in summary if r['success'])
        tot_r = sum(r['rounds'] for r in summary if r['success'])
        log('[{} {} {}] repeat {}: done in {:.1f}s, {}/{} tasks solved, avg {:.1f} '
            'rounds for solved ones'.format(
                job['suite'], job['model'], baseline, rep, time.time() - t0,
                ok, len(summary), (tot_r / ok) if ok else 0.0))
        costs = [r.get('cost_usd') for r in summary]
        cost_txt = ('${:.4f}'.format(sum(c for c in costs if c is not None))
                    + (' (+n/a)' if any(c is None for c in costs) else ''))
        log('LLM calls: {}  |  tokens: {} prompt + {} completion = {} total'.format(
            sum(r['n_calls'] for r in summary),
            sum(r['prompt_tokens'] for r in summary),
            sum(r['completion_tokens'] for r in summary),
            sum(r['total_tokens'] for r in summary)))
        log('LLM time: {:.1f}s  |  verify time: {:.1f}s  |  cost: {}'.format(
            sum(r['llm_time_s'] for r in summary),
            sum(r['verify_time_s'] for r in summary), cost_txt))
        out_json, out_xlsx = next_summary_paths(results_dir, job['model'], baseline)
        with open(out_json, 'w', encoding='utf-8') as f:
            json.dump({'suite': job['suite'], 'model': job['model'],
                       'baseline': baseline, 'max_rounds': max_rounds,
                       'results': summary}, f, ensure_ascii=False, indent=1)
        log('summary written to {}'.format(out_json))
        write_results_xlsx(summary, out_xlsx)
        log('xlsx results written to {}'.format(out_xlsx))
    return 0


def worker_main(args):
    """--worker: run the jobs listed in --jobs (a JSON array) sequentially."""
    jobs = json.loads(args.jobs)
    tag = args.tag or 'worker'
    out_lock = threading.Lock()

    def log(msg):
        with out_lock:
            for line in str(msg).splitlines():
                print('{} | {}'.format(tag, line))
            sys.stdout.flush()

    for i, job in enumerate(jobs, 1):
        log('JOB {}/{}: suite={} model={} baseline={}'.format(
            i, len(jobs), job['suite'], job['model'], job['baseline']))
        try:
            rc = run_job(job, log)
        except KeyboardInterrupt:
            log('interrupted -- stopping this worker')
            return 130
        except Exception as e:
            log('job failed with exception: {}'.format(e))
            rc = 1
        if rc != 0:
            log('job exited with code {} -- continuing with the next job'.format(rc))
    log('worker finished all jobs')
    return 0


# ------------------------------------------------------------- orchestration
def pump(stream, tag, lock):
    """Echo tagged child output while the child runs."""
    for line in iter(stream.readline, ''):
        with lock:
            print('[{}] {}'.format(tag, line.rstrip()))
            sys.stdout.flush()
    stream.close()


def orchestrate(args):
    suites = (['hard', 'complex'] if args.suite == 'both' else [args.suite])
    models = [m.strip() for m in args.models.split(',') if m.strip()]
    baselines = [b.strip() for b in args.baselines.split(',') if b.strip()]
    results_root = os.path.abspath(args.results_root or APP_DIR)

    # -- sanity checks before launching anything ------------------------------
    missing = []
    if not os.path.exists(NUXMV):
        missing.append(NUXMV)
    for s in suites:
        if not os.path.exists(BENCHMARKS[s]):
            missing.append(BENCHMARKS[s])
    if not os.path.exists(PROMPT_FILE):
        missing.append(PROMPT_FILE)
    if missing:
        print('ERROR: missing files -- the distribution folder is incomplete:')
        for m in missing:
            print('  ' + m)
        return 2

    jobs = []
    for suite in suites:
        for model in models:
            for baseline in baselines:
                jobs.append({
                    'suite': suite,
                    'benchmark': BENCHMARKS[suite],
                    'results_dir': os.path.join(results_root,
                                                'results_' + suite),
                    'model': model,
                    'baseline': baseline,
                    'repeat': args.repeat,
                    'max_rounds': args.max_rounds,
                    'tasks': ([t.strip() for t in args.tasks.split(',') if t.strip()]
                              if args.tasks else None),
                    'limit': args.limit,
                    'llm_attempts': args.llm_attempts,
                    'llm_retry_wait': args.llm_retry_wait,
                })

    # -- group jobs into worker processes --------------------------------------
    if args.parallel == 'models':
        groups = [(m, [j for j in jobs if j['model'] == m]) for m in models]
    elif args.parallel == 'baselines':
        groups = [(b, [j for j in jobs if j['baseline'] == b]) for b in baselines]
    else:
        groups = [('sequential', jobs)]

    print('BTRun benchmark driver')
    print('  suites     : {}'.format(' + '.join(suites)))
    print('  models     : {}'.format(', '.join(models)))
    print('  baselines  : {}'.format(', '.join(baselines)))
    print('  repeat     : {}   max_rounds: {}'.format(args.repeat, args.max_rounds))
    print('  parallel   : {} ({} worker processes)'.format(
        args.parallel, len(groups)))
    print('  results in : {}'.format(results_root))
    print('  LLM API    : {} (DIRECT connection, no proxy)'.format(
        load_config()['base_url']))
    total_runs = sum(len(g[1]) for g in groups)
    print('  total jobs : {} (= suites x models x baselines)'.format(total_runs))
    print()

    # -- spawn one child process per group -------------------------------------
    lock = threading.Lock()
    procs = []
    t0 = time.time()
    for tag, group_jobs in groups:
        for j in group_jobs:                      # keep job order visible
            print('[{}] queued: suite={} baseline={} repeat={}'.format(
                tag, j['suite'], j['baseline'], j['repeat']))
        cmd = self_exe_cmd() + ['--worker', '--tag', tag,
                                '--jobs', json.dumps(group_jobs)]
        p = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                             text=True, encoding='utf-8', errors='replace',
                             env=child_env(), bufsize=1)
        th = threading.Thread(target=pump, args=(p.stdout, tag, lock), daemon=True)
        th.start()
        procs.append((tag, p, th))
        print()
    print('all {} worker process(es) started; output below is tagged per worker'
          .format(len(procs)))

    codes = {}
    for tag, p, th in procs:
        p.wait()
        th.join(timeout=30)
        codes[tag] = p.returncode

    print()
    print('================ summary ================')
    for tag, p, th in procs:
        print('worker [{}]: exit code {}'.format(tag, codes[tag]))
    print('all workers finished in {:.1f}s'.format(time.time() - t0))
    print('per-run summaries (json + xlsx) are in:')
    for s in suites:
        print('  ' + os.path.join(results_root, 'results_' + s))
    if any(c != 0 for c in codes.values()):
        return 1
    return 0


# ------------------------------------------------------------------- selftest
def selftest_main(check_api):
    """Check the bundled toolchain: files, NuSMV binary, behaverify verify on
    the reference tree; optionally one live LLM call (direct connection)."""
    print('BTRun selftest -- app dir: {}'.format(APP_DIR))
    ok = True

    def check(name, cond, detail=''):
        nonlocal ok
        print('  [{}] {}{}'.format('OK ' if cond else 'FAIL', name,
                                   (' -- ' + detail) if detail and not cond else ''))
        if not cond:
            ok = False

    check('NuSMV binary', os.path.exists(NUXMV), NUXMV)
    check('hard benchmark xlsx', os.path.exists(BENCHMARKS['hard']))
    check('complex benchmark xlsx', os.path.exists(BENCHMARKS['complex']))
    check('prompt file', os.path.exists(PROMPT_FILE))
    check('selftest tree', os.path.exists(SELFTEST_TREE))

    if os.path.exists(NUXMV) and os.path.exists(SELFTEST_TREE):
        out = os.path.join(APP_DIR, 'selftest', '_out')
        print('  running behaverify+NuSMV on the reference tree (up to {}s)...'
              .format(VERIFY_TIMEOUT))
        t0 = time.time()
        report = verify_tree(SELFTEST_TREE, out)
        print('  verify: status={} ({} true / {} false) in {:.1f}s'.format(
            report['status'], report['n_true'], report['n_false'],
            time.time() - t0))
        check('behaverify+NuSMV verify chain',
              report['status'] == 'pass' and report['n_true'] > 0,
              report.get('message', ''))

    if check_api:
        cfg = load_config()
        check('API key present', bool(cfg.get('api_key')))
        print('  calling {} model {} (direct connection)...'.format(
            cfg['base_url'], cfg['model']))
        t0 = time.time()
        try:
            reply, usage = call_llm(cfg, 'Reply with the single word: OK')
            print('  reply: {}'.format(reply.strip()[:80]))
            print('  usage: {} (in {:.1f}s)'.format(usage, time.time() - t0))
            check('live LLM call', 'ok' in reply.strip().lower()[:20])
        except Exception as e:
            check('live LLM call', False, str(e))

    print('SELFTEST {}'.format('PASSED' if ok else 'FAILED'))
    return 0 if ok else 1


# -------------------------------------------------------------------------- cli
def main():
    # UTF-8 console output on every Windows machine, errors never crash logging
    for s in (sys.stdout, sys.stderr):
        if hasattr(s, 'reconfigure'):
            try:
                s.reconfigure(encoding='utf-8', errors='replace')
            except Exception:
                pass

    # --verify-child must bypass argparse: behaverify flags would be rejected
    if len(sys.argv) > 1 and sys.argv[1] == '--verify-child':
        verify_child_main(sys.argv[2], sys.argv[3], sys.argv[4:])
        return 0

    ap = argparse.ArgumentParser(prog='BTRun',
                                 description='portable LLM->BT benchmark runner')
    ap.add_argument('--suite', choices=['hard', 'complex', 'both'], default='both',
                    help='benchmark suite(s) to run (default: %(default)s)')
    ap.add_argument('--models', default=','.join(MODELS),
                    help='comma separated models (default: %(default)s)')
    ap.add_argument('--baselines', default=','.join(BASELINES),
                    help='comma separated baselines (default: %(default)s)')
    ap.add_argument('--repeat', type=int, default=5,
                    help='runs per (suite, model, baseline) combination (default %(default)s)')
    ap.add_argument('--max_rounds', type=int, default=10,
                    help='LLM call cap per task (default %(default)s)')
    ap.add_argument('--parallel', choices=['models', 'baselines', 'none'],
                    default='models',
                    help='run one worker per model / per baseline / sequentially '
                         '(default: %(default)s)')
    ap.add_argument('--tasks', default=None,
                    help='comma separated task names (filter, for testing)')
    ap.add_argument('--limit', type=int, default=None,
                    help='only run the first N tasks of each suite (for testing)')
    ap.add_argument('--results_root', default=None,
                    help='root folder for results_hard/ and results_complex/ '
                         '(default: folder next to the executable)')
    ap.add_argument('--llm_attempts', type=int, default=LLM_MAX_ATTEMPTS,
                    help='max attempts per LLM call (default %(default)s)')
    ap.add_argument('--llm_retry_wait', type=float, default=LLM_RETRY_WAIT,
                    help='seconds between LLM retries (default %(default)s)')
    ap.add_argument('--worker', action='store_true', help=argparse.SUPPRESS)
    ap.add_argument('--tag', default=None, help=argparse.SUPPRESS)
    ap.add_argument('--jobs', default=None, help=argparse.SUPPRESS)
    ap.add_argument('--selftest', action='store_true',
                    help='check the bundled toolchain (add --api to also test '
                         'the live LLM connection)')
    ap.add_argument('--api', action='store_true',
                    help='with --selftest: also make one live LLM call')
    args = ap.parse_args()

    if args.selftest:
        return selftest_main(args.api)
    if args.worker:
        return worker_main(args)
    return orchestrate(args)


if __name__ == '__main__':
    try:
        sys.exit(main())
    except KeyboardInterrupt:
        print('\ninterrupted by user')
        sys.exit(130)
